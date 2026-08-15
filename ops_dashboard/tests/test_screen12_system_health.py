"""Screen 12 — System Health.

A health screen that shows green for something it did not measure is worse than
no health screen, so most of what follows guards exactly that: UNKNOWN must
never collapse into HEALTHY, an un-instrumented metric must never acquire a
number, and a stale READY must not sit silently beside a dead service.

Plus the approved alert requirement: the alert's own WORDS carry the semantic
colour, and a positive event stays green even inside an Alerts card.

⭐ Assertions are anchored to the fixture's KNOWN values or to a PROPERTY that
must hold. ⛔ No assertion is written by reading back what the code produced.

FIXTURE (conftest), so every number below is traceable:
  preflight run pf_a_1, phase A, overall_status READY, 13 check rows:
    Broker   5 PASS · Database 2 PASS · Engine 3 (2 PASS + capital_deployment WARN)
    State    2 PASS · VM Health 1 PASS
  ⇒ Capital pillar = WARNING (it draws capital_deployment BY NAME)
  autofix log: one SUCCESS, one FAILED, one with NO result ⇒ TRIGGERED
  the trading engine is NOT running under test ⇒ Trading Engine FAILED
  systemctl does not exist on the test host ⇒ every unit UNKNOWN
"""
from __future__ import annotations

import io
import os
import re

from backend.services import system_health as sh


def _read(*parts) -> str:
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(here, *parts), encoding="utf-8") as fh:
        return fh.read()


def _tpl() -> str:
    return _read("frontend", "templates", "services.html")


def _css() -> str:
    return _read("frontend", "static", "style.css")


def _shipped(client) -> str:
    html = client.get("/services").data.decode("utf-8")
    return re.sub(r"<!--.*?-->", " ", html, flags=re.S)


def _h(client, extra=""):
    return client.get("/api/system-health" + extra).get_json()


# ── UNKNOWN IS NOT HEALTHY — the screen's core safety property ───────────────
def test_unknown_outranks_healthy_in_every_rollup():
    """⛔ THE failure this guards: an overall HEALTHY computed over components
    that could not be evaluated. Not knowing is not the same as being well."""
    assert sh._worst(["HEALTHY", "UNKNOWN"]) == "UNKNOWN"
    assert sh._worst(["HEALTHY", "HEALTHY"]) == "HEALTHY"
    assert sh._worst(["UNKNOWN", "WARNING"]) == "WARNING"
    assert sh._worst(["WARNING", "FAILED"]) == "FAILED"
    assert sh._worst(["UNKNOWN", "FAILED"]) == "FAILED"
    # ⛔ an EMPTY set is UNKNOWN, never HEALTHY — nothing measured is not "fine"
    assert sh._worst([]) == "UNKNOWN"


def test_a_host_without_systemctl_reports_unknown_not_failed(client):
    """On a non-Linux host `systemctl` does not exist. ⛔ That must not read as
    every service being down (a false incident) nor as up (a false all-clear)."""
    units = [s for s in _h(client)["services"] if s["kind"] == "systemd"]
    assert units, "the fixture configures systemd units"
    for u in units:
        assert u["status"] == "UNKNOWN", u["service"]
        assert u["raw_state"] == "unavailable"


def test_overall_status_is_never_hard_coded_healthy(client):
    """The trading engine is not running under test, so overall MUST reflect it."""
    d = _h(client)
    assert d["overall"]["status"] == "FAILED"
    engine = [s for s in d["services"] if s["service"] == "Trading Engine"][0]
    assert engine["status"] == "FAILED"
    assert "UNKNOWN outranks HEALTHY" in d["overall"]["rollup"]


def test_counts_partition_the_service_population(client):
    """A PROPERTY: the four counts must always add to the total."""
    d = _h(client)
    assert sum(d["counts"].values()) == d["total_services"] == len(d["services"])


# ── NOT INSTRUMENTED never becomes a number ─────────────────────────────────
def test_cpu_ram_and_network_are_declared_gaps(client):
    """⛔ system_metrics.cpu_pct/memory_mb are -1.0 sentinels (psutil absent) and
    no network metric exists. None of the three may show a value."""
    vm = _h(client)["vm"]
    for key in ("cpu_pct", "ram_pct", "network"):
        assert vm[key]["measured"] is False, key
        assert vm[key]["value"] is None, key
        assert vm[key]["status"] == "UNKNOWN", key
        assert vm[key]["reason"], key
    assert "psutil" in vm["cpu_pct"]["reason"]
    assert "network" in vm["network"]["reason"].lower()


def test_a_gap_is_never_green_on_any_path(client):
    """⛔ Whatever the filter, an un-instrumented metric stays UNKNOWN."""
    for q in ("", "?status=HEALTHY", "?kind=systemd", "?status=FAILED"):
        vm = _h(client, q)["vm"]
        for k, v in vm.items():
            if isinstance(v, dict) and v.get("measured") is False:
                assert v["status"] == "UNKNOWN", (q, k)
                assert v["value"] is None, (q, k)


def test_disk_is_measured_because_it_genuinely_is(client):
    """⭐ The control: disk uses stdlib shutil and works on every platform, so it
    MUST be measured. Without it, "everything is a gap" would pass vacuously."""
    disk = _h(client)["vm"]["disk_pct"]
    assert disk["measured"] is True
    assert isinstance(disk["value"], (int, float)) and 0 <= disk["value"] <= 100


def test_db_connection_count_is_not_applicable_not_zero(client):
    """SQLite is embedded — a count would describe the dashboard, not the trader."""
    db = _h(client)["database"]
    assert db["connection_count"] is None
    assert "no server" in db["connection_count_note"]


def test_trends_declare_which_series_do_not_exist(client):
    tr = _h(client)["trends"]
    for k in ("cpu", "ram", "response_time"):
        assert tr[k]["measured"] is False, k
        assert tr[k]["series"] == [], k
        assert tr[k]["reason"], k
    assert "sentinel" in tr["cpu"]["reason"]


# ── SEMANTIC ALERT COLOUR — the approved requirement ────────────────────────
def test_positive_alerts_are_green_even_on_a_critical_feed():
    """⛔ THE defect the approved design names: a positive event painted red
    merely because it arrived in an Alerts card."""
    for title in ("Connection restored", "Recovery Success",
                  "Service healthy again", "Backup completed successfully",
                  "Broker reconnected"):
        assert sh.classify_alert("CRITICAL", title) == "HEALTHY", title
        assert sh.classify_alert("INFO", title) == "HEALTHY", title


def test_negative_alerts_stay_red():
    for title in ("Backtest Engine is down", "Recovery Failed",
                  "Critical failure in order path", "Broker unreachable"):
        assert sh.classify_alert("CRITICAL", title) == "FAILED", title


def test_a_hopeful_title_cannot_downgrade_a_real_failure():
    """⭐ The override is ONE-WAY. A CRITICAL alert is never demoted below its
    declared severity by a word — only a plain RESOLUTION is promoted."""
    assert sh.classify_alert("CRITICAL", "Order rejected") == "FAILED"
    assert sh.classify_alert("WARNING", "Latency above threshold") == "WARNING"
    # a bare, unmarked title keeps its severity
    assert sh.classify_alert("CRITICAL", "Something happened") == "FAILED"
    assert sh.classify_alert("WARNING", "Something happened") == "WARNING"


def test_severity_maps_to_the_four_statuses():
    assert sh.classify_alert("CRITICAL", "x") == "FAILED"
    assert sh.classify_alert("WARNING", "x") == "WARNING"
    assert sh.classify_alert("INFO", "x") == "UNKNOWN"
    assert sh.classify_alert(None, "x") == "UNKNOWN"


def test_alert_TEXT_carries_the_colour_not_only_an_icon():
    """The approved requirement is explicit: colour the words. The alert title
    element must bind the semantic class, ⛔ not just a dot or a badge."""
    t = _tpl()
    block = t.split('class="sysh-alerts"', 1)[1].split("</ul>", 1)[0]
    assert 'class="sysh-al-t"' in block
    assert ':class="txt(a.status)"' in block
    # and the severity label too
    assert 'class="sysh-sev"' in block


def test_the_css_does_not_repaint_alert_text_over_the_semantic_class():
    """⛔ A colour on `.sysh-al-t` would override the semantic class and make
    every alert one colour — which is the whole defect."""
    css = _css()
    rule = re.search(r"\.sysh-page \.sysh-al-t \{([^}]*)\}", css)
    assert rule, "the rule must exist"
    assert "color" not in rule.group(1)


def test_semantic_colour_helper_maps_all_four_states():
    t = _tpl()
    body = re.search(r"txt\(status\)\s*\{(.*?)\n    \},", t, re.S).group(1)
    assert "HEALTHY: \"num-pos\"" in body
    assert "WARNING: \"num-warn\"" in body
    assert "FAILED: \"num-neg\"" in body
    assert "UNKNOWN" in body and "sysh-unknown" in body


def test_status_colours_are_the_projects_own(client):
    css = _css()
    assert ".sysh-page .sysh-HEALTHY { background: var(--pos-tint);   color: var(--pos); }" in css
    assert ".sysh-page .sysh-WARNING { background: var(--warn-tint);  color: var(--yellow); }" in css
    assert ".sysh-page .sysh-FAILED  { background: var(--neg-tint);   color: var(--neg); }" in css
    assert ".sysh-page .sysh-UNKNOWN { background: var(--muted-tint); color: var(--dim); }" in css


# ── TRADING READINESS ───────────────────────────────────────────────────────
def test_readiness_comes_from_preflight_not_from_overall_status(client):
    """⛔ Readiness is NOT a copy of Overall Status. In this fixture the two
    genuinely DISAGREE — preflight said READY this morning, the engine is down
    now — and that is precisely why they must be separate."""
    d = _h(client)
    assert d["overall"]["status"] == "FAILED"
    assert d["readiness"]["ready"] is True
    assert d["readiness"]["overall_status"] == "READY"


def test_a_stale_ready_beside_a_dead_service_is_flagged(client):
    """⚠️ The screen's headline question is "is trading safe to continue?".
    A READY badge over a failed service must not answer it silently."""
    rd = _h(client)["readiness"]
    c = rd["live_contradiction"]
    assert c is not None
    assert c["severity"] == "FAILED"
    assert "Trading Engine" in c["failed_services"]
    assert "READY" in c["message"]
    # ⛔ and the verdict itself is NOT rewritten — both facts survive
    assert rd["ready"] is True


def test_the_five_pillars_are_evaluated_from_real_checks(client):
    """Broker/Database/Engine/State come from `check_group`; Capital has NO group
    of its own and is assembled from named checks — so it proves both paths."""
    rd = _h(client)["readiness"]
    by = {p["pillar"]: p for p in rd["pillars"]}
    assert set(by) == {"Broker", "Database", "Core Services", "Capital", "Risk"}
    assert by["Broker"]["status"] == "HEALTHY" and by["Broker"]["checks"] == 5
    assert by["Database"]["status"] == "HEALTHY" and by["Database"]["checks"] == 2
    assert by["Risk"]["status"] == "HEALTHY" and by["Risk"]["checks"] == 2
    # capital_deployment is seeded WARN ⇒ Capital is WARNING while others are not
    assert by["Capital"]["status"] == "WARNING"
    assert by["Capital"]["checks"] == 3
    assert by["Capital"]["note"] and "no Capital group" in by["Capital"]["note"]
    assert rd["corroborated"] is True and rd["checks_evaluated"] > 0


def test_readiness_without_a_preflight_run_is_not_ready_by_default(client):
    """⛔ Absence of a verdict is not a READY verdict."""
    out = sh._readiness(None)
    assert out["ready"] is None
    assert out["verdict"] == "UNKNOWN"
    assert "absence of a verdict is not a READY verdict" in out["reason"]
    assert all(p["status"] == "UNKNOWN" for p in out["pillars"])


def test_a_ready_verdict_with_no_checks_is_degraded_not_trusted():
    """⭐ A verdict nothing could have contradicted is not evidence. The system's
    own status is still reported, but the badge is degraded and says why."""
    out = sh._readiness({"overall_status": "READY", "checks": [],
                         "started_at": "2026-08-15T08:30:00", "phase": "A"})
    assert out["ready"] is True                 # ⛔ not rewritten to False
    assert out["verdict"] == "WARNING"          # but not shown as a clean pass
    assert out["corroborated"] is False
    assert "could not be corroborated" in out["reason"]


def test_a_critical_failure_makes_it_not_ready():
    out = sh._readiness({
        "overall_status": "CRITICAL_FAILURE", "phase": "B",
        "started_at": "2026-08-15T09:14:00",
        "checks": [{"check_name": "kite_token_fresh_today", "check_group": "Broker",
                    "criticality": "CRITICAL", "status": "FAIL"}]})
    assert out["ready"] is False
    assert out["verdict"] == "FAILED"
    assert "kite_token_fresh_today" in " ".join(out["blockers"])
    by = {p["pillar"]: p for p in out["pillars"]}
    assert by["Broker"]["status"] == "FAILED"


# ── AUTO-RECOVERY: triggered is not success ─────────────────────────────────
def test_recovery_distinguishes_triggered_success_and_failed(client):
    """⛔ A TRIGGERED attempt must never be counted as a recovery."""
    rec = _h(client)["recovery"]
    assert rec["success"] == 1 and rec["failed"] == 1 and rec["triggered"] == 1
    by = {e["state"]: e for e in rec["events"]}
    assert by["SUCCESS"]["status"] == "HEALTHY"
    assert by["FAILED"]["status"] == "FAILED"
    # ⭐ in-flight is WARNING, ⛔ not green — the component is not fixed yet
    assert by["TRIGGERED"]["status"] == "WARNING"
    assert "never counted as a recovery" in rec["note"]


def test_a_failed_recovery_carries_its_error(client):
    ev = [e for e in _h(client)["recovery"]["events"] if e["state"] == "FAILED"][0]
    assert ev["error"] == "permission denied"


# ── DEPENDENCIES ────────────────────────────────────────────────────────────
def test_the_five_approved_dependencies_are_present(client):
    names = [d["name"] for d in _h(client)["dependencies"]]
    assert names == ["Chartink", "Broker API", "Database", "Tailscale", "Internet"]


def test_a_dependency_is_never_healthy_merely_because_it_is_configured(client):
    """⛔ Internet has no probe anywhere, so it must be UNKNOWN with a reason —
    ⛔ not quietly inherited from the broker check."""
    by = {d["name"]: d for d in _h(client)["dependencies"]}
    assert by["Internet"]["status"] == "UNKNOWN"
    assert by["Internet"]["measured"] is False
    assert "no internet-reachability probe" in by["Internet"]["source"]
    # the one with a real live probe IS measured — the control
    assert by["Database"]["measured"] is True and by["Database"]["status"] == "HEALTHY"


# ── SERVICES / BROKER / UPTIME ──────────────────────────────────────────────
def test_no_fake_services_are_invented(client):
    """The reference's example list names engines that are THREADS inside one
    unit. Only real, sourced services appear."""
    d = _h(client)
    kinds = {s["kind"] for s in d["services"]}
    assert kinds <= {"systemd", "engine", "engine-component", "database", "dashboard"}
    names = {s["service"] for s in d["services"]}
    assert "Database" in names and "Dashboard API" in names and "Trading Engine" in names
    assert "trading-system.service" in names


def test_response_time_is_absent_where_nothing_is_timed(client):
    """⛔ `systemctl` latency is the dashboard's own subprocess cost, not the
    service's responsiveness — so a systemd row carries no response time."""
    for s in _h(client)["services"]:
        if s["kind"] == "systemd":
            assert s["response_ms"] is None
            assert "N/A" in s["response_note"] or "no request" in s["response_note"]
        if s["kind"] == "database":
            assert isinstance(s["response_ms"], (int, float))


def test_broker_login_and_token_are_separate_facts(client):
    """⛔ Collapsing them would hide a token that EXISTS but is STALE — which
    presents as a silent no-trade morning."""
    bk = _h(client)["broker"]
    for k in ("broker_status", "login_status", "token_status"):
        assert bk[k] in ("HEALTHY", "WARNING", "FAILED", "UNKNOWN")
    assert set(bk["token_detail"]) == {"file_exists", "fresh_today"}
    assert bk["last_api_call"]["measured"] is False


def test_uptime_does_not_come_from_the_dashboard(client):
    """⛔ Uptime must not reset when this page refreshes: it is the trading
    engine's own monotonic clock, read from /health."""
    up = _h(client)["uptime"]
    assert "monotonic" in up["source"]
    assert up["measured"] is False and up["seconds"] is None   # engine is down here
    assert "not answering" in up["reason"]


def test_uptime_formatting_never_fabricates():
    assert sh._fmt_uptime(None) is None
    assert sh._fmt_uptime(0) == "0m"
    assert sh._fmt_uptime(90) == "1m"
    assert sh._fmt_uptime(3700) == "1h 01m"
    assert sh._fmt_uptime(7 * 86400 + 14 * 3600 + 32 * 60) == "7d 14h 32m"


def test_systemd_uptime_refuses_an_unusable_stamp():
    """⛔ None, never 0 — a zero uptime reads as "just restarted", which is a
    materially different and alarming statement from "not known"."""
    from backend.readers import host_reader
    assert host_reader._uptime_from_systemd_stamp("") is None
    assert host_reader._uptime_from_systemd_stamp("n/a") is None
    assert host_reader._uptime_from_systemd_stamp("Fri 2026-13-45 99:99:99 IST") is None


# ── FILTERS / VIEW ALL ──────────────────────────────────────────────────────
def test_the_filter_is_a_view_and_the_counts_stay_whole_population(client):
    """This screen is a LIVE SNAPSHOT: filtering the table must not make the KPI
    deck describe a subset, or an operator would lose the failure they filtered
    away from. `services_all` and the counts stay whole-population."""
    a, b = _h(client), _h(client, "?status=UNKNOWN")
    assert len(b["services"]) < len(a["services"])
    assert all(s["status"] == "UNKNOWN" for s in b["services"])
    assert b["counts"] == a["counts"]
    assert b["services_all"] == a["services_all"] == len(a["services"])


def test_kind_filter_narrows_the_table(client):
    d = _h(client, "?kind=systemd")
    assert d["services"] and all(s["kind"] == "systemd" for s in d["services"])


# ── EXPORT ──────────────────────────────────────────────────────────────────
def test_export_carries_every_panel_and_respects_the_view(client):
    from openpyxl import load_workbook
    r = client.get("/api/export/system-health?status=UNKNOWN")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Services", "Readiness", "Dependencies", "Alerts",
                             "Auto-Recovery", "Service Events", "Summary"]
    ws = wb["Services"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Service", "Kind", "Status"]
    body = list(ws.iter_rows(min_row=2, values_only=True))
    assert body and {r[2] for r in body} == {"UNKNOWN"}


def test_export_spells_out_the_gaps(client):
    """⛔ A blank cell would let a reader assume the value was simply zero."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/api/export/system-health").data))
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["CPU %"] == "NOT INSTRUMENTED"
    assert summary["RAM %"] == "NOT INSTRUMENTED"
    assert summary["Network"] == "NOT INSTRUMENTED"
    assert summary["DB connection count"] == "NOT APPLICABLE"
    assert "psutil" in summary["CPU % why"]
    assert summary["Overall Status"] == "FAILED"


def test_export_readiness_sheet_carries_the_pillars(client):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/api/export/system-health").data))
    rows = list(wb["Readiness"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["Broker", "Database", "Core Services",
                                    "Capital", "Risk"]


# ── THE SCREEN ITSELF ───────────────────────────────────────────────────────
def test_route_renders(client):
    r = client.get("/services")
    assert r.status_code == 200
    body = r.data.decode("utf-8")
    assert "sysh-page" in body and "System Health" in body


def test_existing_system_endpoints_are_untouched(client):
    """Screen 12 is ADDITIVE — /api/services and /api/vm keep their contracts."""
    d = client.get("/api/services").get_json()
    assert "units" in d and "cron" in d and "trader" in d
    v = client.get("/api/vm").get_json()
    assert v["cpu_ram_history"]["collected"] is False


def test_real_time_refresh_is_preserved():
    t = _tpl()
    assert '@ops-refresh.window="refresh()"' in t
    body = re.search(r"refresh\(\)\s*\{(.*?)\}", t, re.S).group(1)
    assert "tPage" not in body and "applied" not in body and "cols" not in body


def test_columns_use_the_existing_drag_convention():
    t = _tpl()
    for token in ('draggable="true"', "@dragstart=", "@drop.prevent=", "onDragStart(",
                  "onDrop(", "initCols()", "saveCols()", "resetCols()",
                  "is-drag", "is-over", "screen12.systemhealth.colOrder"):
        assert token in t, token
    body = re.search(r"headClick\(key\)\s*\{(.*?)\n    \},", t, re.S).group(1)
    assert "_dragEndAt" in body and "250" in body


def test_alignment_is_by_column_role_not_blanket():
    css = _css()
    assert ".sysh-page .cap-table th { text-align: center; }" in css
    assert ".sysh-page .cap-table th.lbl { text-align: left; }" in css
    assert ".sysh-page table td { text-align: center" not in css


def test_no_alpine_template_loop_inside_an_svg():
    """⛔ Pinned across the whole tree since Screen 10."""
    import glob
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    offenders = []
    for path in glob.glob(os.path.join(here, "frontend", "templates", "*.html")):
        html = open(path, encoding="utf-8").read()
        for svg in re.findall(r"<svg\b.*?</svg>", html, re.S):
            if re.search(r"<template\b", svg):
                offenders.append(os.path.basename(path))
    assert not offenders, sorted(set(offenders))


def test_no_hard_coded_values_in_the_markup():
    t = _tpl()
    for ghost in ("7d 14h 32m", "18 / 100", "256.4 GB", "120 Mbps", "0.68", "99.82"):
        assert ghost not in t, ghost


def test_css_is_scoped_to_this_screen():
    block = _css().split("SCREEN 12 — SYSTEM HEALTH", 1)[1].split("*/", 1)[1]
    block = re.split(r"SCREEN \d+ [—-]", block, maxsplit=1)[0]
    block = re.sub(r"/\*.*?\*/", " ", block, flags=re.S)
    rules = re.findall(r"([^{}]+)\{([^{}]*)\}", block)
    assert len(rules) > 40
    for selector, _body in rules:
        selector = selector.strip()
        if selector.startswith("@media"):
            continue
        for part in selector.split(","):
            part = part.strip().lstrip("{").strip()
            if part:
                assert part.startswith(".sysh-page"), part


def test_other_screens_are_untouched():
    css = _css()
    for sel in (".cap-page .cap-table th { text-align: center; }",
                ".pnl-page .cap-table th { text-align: center; }",
                ".slp-page .cap-table th { text-align: center; }",
                ".exec-page .cap-table th { text-align: center; }",
                "table td.ctr { text-align: center !important; }"):
        assert sel in css, sel
