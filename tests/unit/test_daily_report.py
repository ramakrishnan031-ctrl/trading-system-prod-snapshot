"""
tests/unit/test_daily_report.py — Trading System v2

Unit tests for reports/daily_report.py (7-sheet xlsx generator).
Target: 20+ tests covering all sheet builders and helpers.
"""

from __future__ import annotations

import json
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List
from unittest.mock import MagicMock, patch

import pytest

from reports.daily_report import (
    ReportData,
    _fmt_time,
    _fmt_datetime,
    _calc_slip_pct,
    _generate_tune_suggestions,
    _get_order_for_trade_leg,
    is_holiday_or_weekend,
    load_report_data,
    generate_daily_report,
    build_sheet_0_dashboard,
    build_sheet_1_signals,
    build_sheet_2_orders,
    build_sheet_3_capital,
    build_sheet_4_candles,
    build_sheet_5_telegram,
    build_sheet_6_strategy,
)


# ─────────────────────────────────────────────────────────────────────────────
# Fixtures
# ─────────────────────────────────────────────────────────────────────────────

@pytest.fixture
def sample_trade():
    """A sample closed trade dict."""
    return {
        "trade_id": "trade-001",
        "signal_id": "sig-001",
        "symbol": "RELIANCE",
        "direction": "LONG",
        "strategy": "MOMENTUM",
        "qty_planned": 10,
        "qty_filled": 10,
        "entry_target_price": 2500.0,
        "entry_actual_price": 2502.0,
        "sl_initial": 2450.0,
        "tgt_initial": 2600.0,
        "margin_reserved": 5000.0,
        "created_at": "2026-05-15T09:30:00+05:30",
        "entry_time": "2026-05-15T09:31:00+05:30",
        "exit_time": "2026-05-15T11:00:00+05:30",
        "exit_price": 2600.0,
        "exit_reason": "TGT_HIT",
        "gross_pnl": 980.0,
        "charges": 50.0,
        "net_pnl": 930.0,
        "status": "CLOSED",
    }


@pytest.fixture
def sample_signal():
    """A sample signal dict."""
    return {
        "signal_id": "sig-001",
        "symbol": "RELIANCE",
        "scanner": "MOMENTUM_BREAK",
        "strategy": "MOMENTUM",
        "triggered_at": "2026-05-15T09:29:00+05:30",
        "received_at": "2026-05-15T09:29:30+05:30",
        "status": "TRADED",
        "trade_id": "trade-001",
        "trigger_price": 2498.0,
    }


@pytest.fixture
def sample_order():
    """A sample order dict."""
    return {
        "order_id": "broker-12345",
        "trade_id": "trade-001",
        "leg": "ENTRY",
        "transaction_type": "BUY",
        "order_type": "LIMIT",
        "product": "MIS",
        "variety": "regular",
        "qty_requested": 10,
        "price": 2500.0,
        "trigger_price": None,
        "status": "COMPLETE",
        "qty_filled": 10,
        "avg_fill_price": 2502.0,
        "placed_at": "2026-05-15T09:30:00+05:30",
    }


@pytest.fixture
def sample_report_data(sample_trade, sample_signal, sample_order):
    """Complete ReportData for testing sheet builders."""
    return ReportData(
        date_iso="2026-05-15",
        mode="PAPER",
        account="TEST001",
        opening_capital=100000.0,
        closing_capital_broker=100930.0,
        signals=[sample_signal],
        trades=[sample_trade],
        orders=[sample_order],
        fm_ledger=[
            {
                "ledger_id": 1,
                "ts": "2026-05-15T09:00:00+05:30",
                "entry_type": "INIT",
                "amount": 100000.0,
                "balance_before": 0.0,
                "balance_after": 100000.0,
            }
        ],
        screener_results=[
            {
                "signal_id": "sig-001",
                "score": 75,
                "tier": "HIGH",
                "status": "PASSED",
                "step_results": "{}",
                "latencies": "{}",
                "ts": "2026-05-15T09:29:35+05:30",
            }
        ],
        innings=[],
        system_events=[
            {
                "event_id": 1,
                "timestamp": "2026-05-15T09:15:00+05:30",
                "event_type": "STARTUP",
                "scenario": "COLD",
            }
        ],
        recon_log=[],
        gate_state=[],
        excluded_symbols=["E2E"],
    )


# ─────────────────────────────────────────────────────────────────────────────
# Helper function tests
# ─────────────────────────────────────────────────────────────────────────────

class TestFormatHelpers:
    """Tests for formatting helper functions."""

    def test_fmt_time_valid_iso(self):
        result = _fmt_time("2026-05-15T09:30:45+05:30")
        assert result == "09:30:45"

    def test_fmt_time_empty(self):
        assert _fmt_time("") == ""
        assert _fmt_time(None) == ""

    def test_fmt_time_invalid(self):
        result = _fmt_time("not-a-timestamp")
        assert result == "not-a-ti"

    def test_fmt_datetime_valid_iso(self):
        result = _fmt_datetime("2026-05-15T09:30:45+05:30")
        assert result == "2026-05-15 09:30:45"

    def test_fmt_datetime_empty(self):
        assert _fmt_datetime("") == ""
        assert _fmt_datetime(None) == ""

    def test_calc_slip_pct_positive(self):
        pct = _calc_slip_pct(100.0, 102.0)
        assert pct == 2.0

    def test_calc_slip_pct_zero_base(self):
        assert _calc_slip_pct(0.0, 102.0) == 0.0


class TestOrderLookup:
    """Tests for order lookup helper."""

    def test_get_order_for_trade_leg_found(self, sample_order):
        orders = [sample_order]
        result = _get_order_for_trade_leg(orders, "trade-001", "ENTRY")
        assert result is not None
        assert result["order_id"] == "broker-12345"

    def test_get_order_for_trade_leg_not_found(self, sample_order):
        orders = [sample_order]
        result = _get_order_for_trade_leg(orders, "trade-001", "SL")
        assert result is None

    def test_get_order_for_trade_leg_wrong_trade(self, sample_order):
        orders = [sample_order]
        result = _get_order_for_trade_leg(orders, "trade-999", "ENTRY")
        assert result is None


# ─────────────────────────────────────────────────────────────────────────────
# Holiday/weekend tests
# ─────────────────────────────────────────────────────────────────────────────

class TestHolidayCheck:
    """Tests for holiday/weekend checking."""

    def test_weekend_saturday(self, tmp_path):
        # 2026-05-16 is Saturday
        assert is_holiday_or_weekend("2026-05-16", tmp_path) is True

    def test_weekend_sunday(self, tmp_path):
        # 2026-05-17 is Sunday
        assert is_holiday_or_weekend("2026-05-17", tmp_path) is True

    def test_weekday_no_holiday_file(self, tmp_path):
        # 2026-05-15 is Friday
        assert is_holiday_or_weekend("2026-05-15", tmp_path) is False

    def test_weekday_with_holiday(self, tmp_path):
        import yaml
        holiday_file = tmp_path / "nse_holidays_2026.yaml"
        holiday_file.write_text(yaml.dump({"holidays": ["2026-05-15"]}))

        assert is_holiday_or_weekend("2026-05-15", tmp_path) is True

    def test_weekday_not_in_holiday_list(self, tmp_path):
        import yaml
        holiday_file = tmp_path / "nse_holidays_2026.yaml"
        holiday_file.write_text(yaml.dump({"holidays": ["2026-01-26"]}))

        assert is_holiday_or_weekend("2026-05-15", tmp_path) is False


# ─────────────────────────────────────────────────────────────────────────────
# Tune suggestion tests
# ─────────────────────────────────────────────────────────────────────────────

class TestTuneSuggestions:
    """Tests for auto-tune suggestion generation."""

    def test_tgt_hit_generates_positive_suggestion(self, sample_report_data):
        suggestions = _generate_tune_suggestions(sample_report_data)
        assert any("TGT hit perfectly" in s for s in suggestions)

    def test_sl_hit_adverse_move_suggests_widen(self):
        data = ReportData(
            date_iso="2026-05-15",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=99000.0,
            signals=[],
            trades=[{
                "trade_id": "t1",
                "symbol": "INFY",
                "direction": "LONG",
                "sl_initial": 1500.0,
                "tgt_initial": 1600.0,
                "entry_actual_price": 1550.0,
                "exit_price": 1492.0,  # 0.5% below SL
                "exit_reason": "SL_HIT",
                "status": "CLOSED",
                "net_pnl": -580.0,
                "gross_pnl": -580.0,
                "strategy": "TEST",
            }],
            orders=[],
            fm_ledger=[],
            screener_results=[],
            innings=[],
            system_events=[],
            recon_log=[],
            gate_state=[],
            excluded_symbols=[],
        )
        suggestions = _generate_tune_suggestions(data)
        assert any("widening SL" in s for s in suggestions)

    def test_high_entry_slippage_generates_critical(self):
        data = ReportData(
            date_iso="2026-05-15",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=99000.0,
            signals=[],
            trades=[{
                "trade_id": "t1",
                "symbol": "TCS",
                "direction": "LONG",
                "entry_target_price": 3500.0,
                "entry_actual_price": 3580.0,  # 2.3% slip
                "sl_initial": 3400.0,
                "tgt_initial": 3600.0,
                "exit_price": 3400.0,
                "exit_reason": "SL_HIT",
                "status": "CLOSED",
                "net_pnl": -200.0,
                "gross_pnl": -200.0,
                "strategy": "TEST",
            }],
            orders=[],
            fm_ledger=[],
            screener_results=[],
            innings=[],
            system_events=[],
            recon_log=[],
            gate_state=[],
            excluded_symbols=[],
        )
        suggestions = _generate_tune_suggestions(data)
        assert any("CRITICAL" in s and "slippage" in s for s in suggestions)

    def test_general_summary_always_generated(self, sample_report_data):
        suggestions = _generate_tune_suggestions(sample_report_data)
        assert any("GENERAL:" in s for s in suggestions)


# ─────────────────────────────────────────────────────────────────────────────
# Sheet builder tests
# ─────────────────────────────────────────────────────────────────────────────

class TestSheetBuilders:
    """Tests for individual sheet builder functions."""

    def test_build_sheet_0_dashboard_creates_sheet(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_0_dashboard(wb, sample_report_data)

        assert ws.title == "0_EOD_Dashboard"
        assert "EOD Dashboard" in str(ws.cell(row=1, column=1).value)

    def test_build_sheet_0_includes_all_sections(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_0_dashboard(wb, sample_report_data)

        cell_values = [ws.cell(row=r, column=1).value for r in range(1, 50)]
        text = " ".join(str(v) for v in cell_values if v)

        assert "Day Overview" in text
        assert "Signal Funnel" in text
        assert "P&L Summary" in text
        assert "System Health" in text
        assert "Tuning" in text

    def test_build_sheet_1_signals_creates_sheet(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_1_signals(wb, sample_report_data)

        assert ws.title == "1_Signals"
        assert ws.cell(row=1, column=1).value == "Trading Date"

    def test_build_sheet_1_signals_includes_recon_columns(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_1_signals(wb, sample_report_data)

        headers = [ws.cell(row=1, column=c).value for c in range(1, 23)]
        assert "Delta" in headers
        assert "Total Rcvd" in headers

    def test_build_sheet_2_orders_creates_sheet(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_2_orders(wb, sample_report_data)

        assert ws.title == "2_Orders"

    def test_build_sheet_3_capital_has_opening_row(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_3_capital(wb, sample_report_data)

        assert ws.title == "3_Capital"
        assert ws.cell(row=2, column=3).value == "Opening"

    def test_build_sheet_3_capital_has_reconciliation(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_3_capital(wb, sample_report_data)

        all_values = []
        for row in ws.iter_rows(min_row=1, max_row=30, max_col=2, values_only=True):
            all_values.extend([str(v) for v in row if v])

        text = " ".join(all_values)
        assert "RECONCILIATION" in text

    def test_build_sheet_3_capital_strategy_fallback(self, sample_report_data):
        """Strategy column should fallback to signal.strategy if trade.strategy is empty."""
        import openpyxl

        # Create a trade with empty strategy
        trade_no_strategy = {
            "trade_id": "trade-002",
            "signal_id": "sig-001",  # Same signal as sample_signal which has strategy="MOMENTUM"
            "symbol": "TCS",
            "direction": "LONG",
            "strategy": "",  # Empty strategy
            "qty_planned": 5,
            "qty_filled": 5,
            "entry_target_price": 3500.0,
            "entry_actual_price": 3505.0,
            "sl_initial": 3450.0,
            "tgt_initial": 3600.0,
            "margin_reserved": 3500.0,
            "created_at": "2026-05-15T09:35:00+05:30",
            "entry_time": "2026-05-15T09:36:00+05:30",
            "exit_time": "2026-05-15T10:00:00+05:30",
            "exit_price": 3600.0,
            "exit_reason": "TGT_HIT",
            "gross_pnl": 475.0,
            "charges": 25.0,
            "net_pnl": 450.0,
            "status": "CLOSED",
        }

        data_with_empty_strategy = ReportData(
            date_iso=sample_report_data.date_iso,
            mode=sample_report_data.mode,
            account=sample_report_data.account,
            opening_capital=sample_report_data.opening_capital,
            closing_capital_broker=sample_report_data.closing_capital_broker,
            signals=sample_report_data.signals,  # Contains sig-001 with strategy="MOMENTUM"
            trades=[trade_no_strategy],
            orders=sample_report_data.orders,
            fm_ledger=sample_report_data.fm_ledger,
            screener_results=sample_report_data.screener_results,
            innings=sample_report_data.innings,
            system_events=sample_report_data.system_events,
            recon_log=sample_report_data.recon_log,
            gate_state=sample_report_data.gate_state,
            excluded_symbols=sample_report_data.excluded_symbols,
        )

        wb = openpyxl.Workbook()
        ws = build_sheet_3_capital(wb, data_with_empty_strategy)

        # Check that row 3 (first trade row after opening) has strategy from signal
        # Column 4 is Strategy
        strategy_value = ws.cell(row=3, column=4).value
        assert strategy_value == "MOMENTUM", f"Expected 'MOMENTUM' from signal fallback, got {strategy_value!r}"

    def test_build_sheet_2_orders_strategy_fallback_to_signal(self, sample_report_data):
        """Strategy col (C) in Sheet 2_Orders should fallback to signal.strategy when trade.strategy is empty."""
        import openpyxl

        trade_no_strategy = {
            "trade_id": "trade-004",
            "signal_id": "sig-001",  # Signal has strategy="MOMENTUM"
            "symbol": "TCS",
            "direction": "LONG",
            "strategy": "",
            "qty_planned": 5,
            "qty_filled": 5,
            "entry_target_price": 3500.0,
            "entry_actual_price": 3505.0,
            "sl_initial": 3450.0,
            "tgt_initial": 3600.0,
            "margin_reserved": 3500.0,
            "created_at": "2026-05-15T09:35:00+05:30",
            "entry_time": "2026-05-15T09:36:00+05:30",
            "exit_time": "2026-05-15T10:00:00+05:30",
            "exit_price": 3600.0,
            "exit_reason": "TGT_HIT",
            "gross_pnl": 475.0,
            "charges": 25.0,
            "net_pnl": 450.0,
            "status": "CLOSED",
        }

        data = ReportData(
            date_iso=sample_report_data.date_iso,
            mode=sample_report_data.mode,
            account=sample_report_data.account,
            opening_capital=sample_report_data.opening_capital,
            closing_capital_broker=sample_report_data.closing_capital_broker,
            signals=sample_report_data.signals,
            trades=[trade_no_strategy],
            orders=sample_report_data.orders,
            fm_ledger=sample_report_data.fm_ledger,
            screener_results=sample_report_data.screener_results,
            innings=sample_report_data.innings,
            system_events=sample_report_data.system_events,
            recon_log=sample_report_data.recon_log,
            gate_state=sample_report_data.gate_state,
            excluded_symbols=sample_report_data.excluded_symbols,
        )

        wb = openpyxl.Workbook()
        ws = build_sheet_2_orders(wb, data)

        # Row 4 = first trade row (rows 1-3 are headers); col 3 = Strategy
        strategy_value = ws.cell(row=4, column=3).value
        assert strategy_value == "MOMENTUM", f"Expected 'MOMENTUM' from signal fallback, got {strategy_value!r}"

    def test_build_sheet_4_candles_strategy_fallback_to_signal(self, sample_report_data):
        """Strategy col (C) in Sheet 4_Candles should fallback to signal.strategy when trade.strategy is empty."""
        import openpyxl

        trade_no_strategy = {
            "trade_id": "trade-005",
            "signal_id": "sig-001",  # Signal has strategy="MOMENTUM"
            "symbol": "INFY",
            "direction": "LONG",
            "strategy": "",
            "qty_planned": 8,
            "qty_filled": 8,
            "entry_target_price": 1500.0,
            "entry_actual_price": 1502.0,
            "sl_initial": 1470.0,
            "tgt_initial": 1560.0,
            "margin_reserved": 2400.0,
            "created_at": "2026-05-15T10:00:00+05:30",
            "entry_time": "2026-05-15T10:01:00+05:30",
            "exit_time": "2026-05-15T11:30:00+05:30",
            "exit_price": 1560.0,
            "exit_reason": "TGT_HIT",
            "gross_pnl": 464.0,
            "charges": 24.0,
            "net_pnl": 440.0,
            "status": "CLOSED",
        }

        data = ReportData(
            date_iso=sample_report_data.date_iso,
            mode=sample_report_data.mode,
            account=sample_report_data.account,
            opening_capital=sample_report_data.opening_capital,
            closing_capital_broker=sample_report_data.closing_capital_broker,
            signals=sample_report_data.signals,
            trades=[trade_no_strategy],
            orders=sample_report_data.orders,
            fm_ledger=sample_report_data.fm_ledger,
            screener_results=sample_report_data.screener_results,
            innings=sample_report_data.innings,
            system_events=sample_report_data.system_events,
            recon_log=sample_report_data.recon_log,
            gate_state=sample_report_data.gate_state,
            excluded_symbols=sample_report_data.excluded_symbols,
        )

        wb = openpyxl.Workbook()
        ws = build_sheet_4_candles(wb, data)

        # Row 3 = first trade row; col 3 = Strategy
        strategy_value = ws.cell(row=3, column=3).value
        assert strategy_value == "MOMENTUM", f"Expected 'MOMENTUM' from signal fallback, got {strategy_value!r}"

    def test_build_sheet_4_candles_creates_sheet(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_4_candles(wb, sample_report_data)

        assert ws.title == "4_Candles"

    def test_build_sheet_4_candles_excludes_cancelled_trades(self, sample_report_data):
        """Candle analysis should exclude CANCELLED trades to avoid NaN values."""
        import openpyxl

        # Add a CANCELLED trade to the data
        cancelled_trade = {
            "trade_id": "trade-cancelled",
            "signal_id": "sig-002",
            "symbol": "INFY",
            "direction": "LONG",
            "strategy": "MOMENTUM",
            "status": "CANCELLED",
            # Fields below would be None/0 for cancelled trades
            "entry_actual_price": None,
            "exit_price": None,
            "exit_reason": None,
        }

        data_with_cancelled = ReportData(
            date_iso=sample_report_data.date_iso,
            mode=sample_report_data.mode,
            account=sample_report_data.account,
            opening_capital=sample_report_data.opening_capital,
            closing_capital_broker=sample_report_data.closing_capital_broker,
            signals=sample_report_data.signals,
            trades=[sample_report_data.trades[0], cancelled_trade],  # 1 CLOSED + 1 CANCELLED
            orders=sample_report_data.orders,
            fm_ledger=sample_report_data.fm_ledger,
            screener_results=sample_report_data.screener_results,
            innings=sample_report_data.innings,
            system_events=sample_report_data.system_events,
            recon_log=sample_report_data.recon_log,
            gate_state=sample_report_data.gate_state,
            excluded_symbols=sample_report_data.excluded_symbols,
        )

        wb = openpyxl.Workbook()
        ws = build_sheet_4_candles(wb, data_with_cancelled)

        # Count trade rows (skip headers in rows 1-2, check from row 3 onwards)
        # Only CLOSED trades should appear, not CANCELLED
        trade_rows = []
        symbol_col = 4  # Symbol column
        for row in range(3, 50):  # Check up to row 50
            symbol = ws.cell(row=row, column=symbol_col).value
            if symbol and symbol != "":
                trade_rows.append((row, symbol))

        # Should have exactly 1 trade (the CLOSED one)
        assert len(trade_rows) == 1, f"Expected 1 CLOSED trade, found {len(trade_rows)}: {trade_rows}"

        # Verify the CLOSED trade is present (RELIANCE from sample_trade)
        assert trade_rows[0][1] == "RELIANCE"

    def test_build_sheet_5_telegram_creates_sheet(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_5_telegram(wb, sample_report_data)

        assert ws.title == "5_Telegram"

    def test_build_sheet_6_strategy_has_two_tables(self, sample_report_data):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_6_strategy(wb, sample_report_data)

        assert ws.title == "6_Strategy_Analysis"

        all_values = []
        for row in ws.iter_rows(min_row=1, max_row=30, max_col=1, values_only=True):
            all_values.extend([str(v) for v in row if v])

        text = " ".join(all_values)
        assert "STRATEGY-WISE" in text
        assert "TIME-OF-DAY" in text

    def test_build_sheet_6_strategy_fallback_to_signal(self, sample_report_data):
        """Strategy analysis should use signal.strategy if trade.strategy is empty."""
        import openpyxl

        # Create a trade with empty strategy
        trade_no_strategy = {
            "trade_id": "trade-003",
            "signal_id": "sig-001",  # Same signal which has strategy="MOMENTUM"
            "symbol": "INFY",
            "direction": "SHORT",
            "strategy": "",  # Empty strategy
            "qty_planned": 8,
            "qty_filled": 8,
            "entry_target_price": 1500.0,
            "entry_actual_price": 1498.0,
            "sl_initial": 1550.0,
            "tgt_initial": 1450.0,
            "margin_reserved": 2400.0,
            "created_at": "2026-05-15T10:00:00+05:30",
            "entry_time": "2026-05-15T10:01:00+05:30",
            "exit_time": "2026-05-15T11:30:00+05:30",
            "exit_price": 1450.0,
            "exit_reason": "TGT_HIT",
            "gross_pnl": 384.0,
            "charges": 20.0,
            "net_pnl": 364.0,
            "status": "CLOSED",
        }

        data_with_empty_strategy = ReportData(
            date_iso=sample_report_data.date_iso,
            mode=sample_report_data.mode,
            account=sample_report_data.account,
            opening_capital=sample_report_data.opening_capital,
            closing_capital_broker=sample_report_data.closing_capital_broker,
            signals=sample_report_data.signals,  # Contains sig-001 with strategy="MOMENTUM"
            trades=[trade_no_strategy],
            orders=sample_report_data.orders,
            fm_ledger=sample_report_data.fm_ledger,
            screener_results=sample_report_data.screener_results,
            innings=sample_report_data.innings,
            system_events=sample_report_data.system_events,
            recon_log=sample_report_data.recon_log,
            gate_state=sample_report_data.gate_state,
            excluded_symbols=sample_report_data.excluded_symbols,
        )

        wb = openpyxl.Workbook()
        ws = build_sheet_6_strategy(wb, data_with_empty_strategy)

        # Check that row 3 (first strategy row) has "MOMENTUM" from signal fallback
        # Column 2 is Strategy
        strategy_value = ws.cell(row=3, column=2).value
        assert strategy_value == "MOMENTUM", f"Expected 'MOMENTUM' from signal fallback, got {strategy_value!r}"

    def test_build_sheet_4_candles_ohlc_populated_from_db(self, sample_report_data):
        """When candle_map has data, OHLC cols H-K should be populated."""
        import openpyxl
        from dataclasses import replace

        data = replace(sample_report_data, candle_map={
            ("RELIANCE", "09:31"): {
                "open": 2490.0, "high": 2610.0,
                "low": 2488.0, "close": 2605.0, "is_synthetic": 0,
            },
        })
        wb = openpyxl.Workbook()
        ws = build_sheet_4_candles(wb, data)

        assert ws.cell(row=3, column=8).value == 2490.0, "Open mismatch"
        assert ws.cell(row=3, column=9).value == 2610.0, "High mismatch"
        assert ws.cell(row=3, column=10).value == 2488.0, "Low mismatch"
        assert ws.cell(row=3, column=11).value == 2605.0, "Close mismatch"
        assert ws.cell(row=3, column=12).value == "No"

    def test_build_sheet_4_candles_ohlc_blank_without_candle_data(self, sample_report_data):
        """Without candle data, OHLC cols and Synthetic? col should stay blank."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_4_candles(wb, sample_report_data)

        assert ws.cell(row=3, column=8).value == ""
        assert ws.cell(row=3, column=9).value == ""
        assert ws.cell(row=3, column=10).value == ""
        assert ws.cell(row=3, column=11).value == ""
        assert ws.cell(row=3, column=12).value == ""

    def test_build_sheet_6_drawdown_pct_positive_trade(self, sample_report_data):
        """When only wins exist, max_loss=0 so drawdown_pct=0."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = build_sheet_6_strategy(wb, sample_report_data)

        # Row 3 = first strategy; col 17 = Drawdown %
        # sample_trade net_pnl=930 (win only), max_loss=0, drawdown=0
        drawdown = ws.cell(row=3, column=17).value
        assert isinstance(drawdown, float), f"Expected float, got {drawdown!r}"
        assert drawdown == 0.0

    def test_build_sheet_6_drawdown_pct_loss_trade(self):
        """Drawdown % = (max_loss / capital_used) * 100."""
        import openpyxl

        data = ReportData(
            date_iso="2026-05-15",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=99000.0,
            signals=[],
            trades=[{
                "trade_id": "t1",
                "symbol": "INFY",
                "direction": "LONG",
                "strategy": "TEST",
                "net_pnl": -500.0,
                "gross_pnl": -500.0,
                "charges": 0,
                "margin_reserved": 10000.0,
                "status": "CLOSED",
            }],
            orders=[], fm_ledger=[], screener_results=[], innings=[],
            system_events=[], recon_log=[], gate_state=[],
            excluded_symbols=[],
        )

        wb = openpyxl.Workbook()
        ws = build_sheet_6_strategy(wb, data)

        # max_loss=-500, capital_used=10000 → -500/10000*100 = -5.0
        drawdown = ws.cell(row=3, column=17).value
        assert drawdown == -5.0, f"Expected -5.0, got {drawdown!r}"

    def test_build_sheet_6_drawdown_pct_zero_capital_guard(self):
        """Drawdown % should be 0.0 when capital_used is 0 (div/zero guard)."""
        import openpyxl

        data = ReportData(
            date_iso="2026-05-15",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=99000.0,
            signals=[],
            trades=[{
                "trade_id": "t1",
                "symbol": "TCS",
                "direction": "LONG",
                "strategy": "TEST",
                "net_pnl": -100.0,
                "gross_pnl": -100.0,
                "charges": 0,
                "margin_reserved": 0,  # zero capital — guard against div/zero
                "status": "CLOSED",
            }],
            orders=[], fm_ledger=[], screener_results=[], innings=[],
            system_events=[], recon_log=[], gate_state=[],
            excluded_symbols=[],
        )

        wb = openpyxl.Workbook()
        ws = build_sheet_6_strategy(wb, data)

        drawdown = ws.cell(row=3, column=17).value
        assert drawdown == 0.0, f"Expected 0.0 for zero capital, got {drawdown!r}"


# ─────────────────────────────────────────────────────────────────────────────
# Integration test
# ─────────────────────────────────────────────────────────────────────────────

class TestGenerateReport:
    """Integration test for full report generation."""

    def _setup_mock_store(self, sample_report_data):
        mock_store = MagicMock()
        mock_store.get_signals_for_date.return_value = sample_report_data.signals
        mock_store.get_trades_for_date.return_value = sample_report_data.trades
        mock_store.get_orders_for_date.return_value = sample_report_data.orders
        mock_store.get_fm_ledger_for_date.return_value = sample_report_data.fm_ledger
        mock_store.get_screener_results_for_date.return_value = sample_report_data.screener_results
        mock_store.get_innings_for_date.return_value = []
        mock_store.get_system_events_for_date.return_value = sample_report_data.system_events
        mock_store.get_reconciliation_log_for_date.return_value = []
        mock_store.get_all_gate_state.return_value = []
        mock_store.get_candles_for_date.return_value = []
        mock_store.get_trade_excursions_for_date.return_value = []
        mock_store.get_session_row.return_value = {
            "mode": "PAPER",
            "account_id": "TEST001",
        }
        return mock_store

    def test_generate_daily_report_creates_file(self, sample_report_data, tmp_path):
        mock_store = self._setup_mock_store(sample_report_data)

        config_dir = tmp_path / "config"
        config_dir.mkdir()
        (config_dir / "system_config.yaml").write_text("excluded_symbols: []\n")

        output_dir = tmp_path / "output"

        result = generate_daily_report(
            store=mock_store,
            date_iso="2026-05-15",
            output_dir=output_dir,
            config_dir=config_dir,
        )

        assert result.exists()
        assert result.name == "daily_report_2026-05-15.xlsx"

    def test_generate_daily_report_has_all_sheets(self, sample_report_data, tmp_path):
        import openpyxl

        mock_store = self._setup_mock_store(sample_report_data)

        config_dir = tmp_path / "config"
        config_dir.mkdir()
        (config_dir / "system_config.yaml").write_text("excluded_symbols: []\n")

        output_dir = tmp_path / "output"

        result = generate_daily_report(
            store=mock_store,
            date_iso="2026-05-15",
            output_dir=output_dir,
            config_dir=config_dir,
        )

        wb = openpyxl.load_workbook(result)
        sheet_names = wb.sheetnames

        assert "0_EOD_Dashboard" in sheet_names
        assert "1_Signals" in sheet_names
        assert "2_Orders" in sheet_names
        assert "3_Capital" in sheet_names
        assert "4_Candles" in sheet_names
        assert "5_Telegram" in sheet_names
        assert "6_Strategy_Analysis" in sheet_names
        assert len(sheet_names) == 7


# ─────────────────────────────────────────────────────────────────────────────
# Edge cases
# ─────────────────────────────────────────────────────────────────────────────

class TestEdgeCases:
    """Edge case tests."""

    def test_empty_data_generates_report(self, tmp_path):
        import openpyxl

        data = ReportData(
            date_iso="2026-05-15",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=100000.0,
            signals=[],
            trades=[],
            orders=[],
            fm_ledger=[],
            screener_results=[],
            innings=[],
            system_events=[],
            recon_log=[],
            gate_state=[],
            excluded_symbols=[],
        )

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        build_sheet_0_dashboard(wb, data)
        build_sheet_1_signals(wb, data)
        build_sheet_2_orders(wb, data)
        build_sheet_3_capital(wb, data)
        build_sheet_4_candles(wb, data)
        build_sheet_5_telegram(wb, data)
        build_sheet_6_strategy(wb, data)

        output_path = tmp_path / "test_report.xlsx"
        wb.save(output_path)

        assert output_path.exists()

    def test_missing_optional_fields_handled(self):
        trade = {
            "trade_id": "t1",
            "symbol": "TEST",
            "status": "CLOSED",
        }
        data = ReportData(
            date_iso="2026-05-15",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=100000.0,
            signals=[],
            trades=[trade],
            orders=[],
            fm_ledger=[],
            screener_results=[],
            innings=[],
            system_events=[],
            recon_log=[],
            gate_state=[],
            excluded_symbols=[],
        )

        suggestions = _generate_tune_suggestions(data)
        assert isinstance(suggestions, list)


# ─────────────────────────────────────────────────────────────────────────────
# Per-strategy effective min_score tests (FIX-119)
# ─────────────────────────────────────────────────────────────────────────────

class TestPerStrategyEligibleScore:
    """Col N (Sheet 1) and Col F (Sheet 2) read eligible_score from screener_results."""

    def _make_data(self, eligible_score):
        signal = {
            "signal_id": "sig-A",
            "symbol": "TATAMOTORS",
            "strategy": "gap_fade_long",
            "scanner": "chartink",
            "status": "TRADED",
            "trade_id": "tr-A",
            "received_at": "2026-05-18T09:30:00+05:30",
            "triggered_at": "2026-05-18T09:29:55+05:30",
            "trigger_price": 800.0,
        }
        trade = {
            "trade_id": "tr-A",
            "signal_id": "sig-A",
            "symbol": "TATAMOTORS",
            "direction": "LONG",
            "strategy": "gap_fade_long",
            "qty_planned": 5,
            "qty_filled": 5,
            "entry_target_price": 800.0,
            "entry_actual_price": 801.0,
            "sl_initial": 780.0,
            "tgt_initial": 840.0,
            "gross_pnl": 500.0,
            "charges": 20.0,
            "net_pnl": 480.0,
            "status": "CLOSED",
            "created_at": "2026-05-18T09:30:00+05:30",
            "entry_time": "2026-05-18T09:30:05+05:30",
            "exit_time": "2026-05-18T10:45:00+05:30",
        }
        screener = {"signal_id": "sig-A", "score": 35, "status": "PASSED"}
        if eligible_score is not None:
            screener["eligible_score"] = eligible_score
        return ReportData(
            date_iso="2026-05-18",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=100480.0,
            signals=[signal],
            trades=[trade],
            orders=[],
            fm_ledger=[],
            screener_results=[screener],
            innings=[],
            system_events=[],
            recon_log=[],
            gate_state=[],
            excluded_symbols=[],
        )

    def test_sheet_1_signals_eligible_score_from_db(self):
        import openpyxl
        data = self._make_data(30)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        build_sheet_1_signals(wb, data)
        ws = wb["1_Signals"]
        eligible_score_cell = ws.cell(row=2, column=14).value
        assert eligible_score_cell == 30, f"Expected 30, got {eligible_score_cell}"

    def test_sheet_2_orders_eligible_score_from_db(self):
        import openpyxl
        data = self._make_data(30)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        build_sheet_2_orders(wb, data)
        ws = wb["2_Orders"]
        eligible_score_cell = ws.cell(row=4, column=6).value
        assert eligible_score_cell == 30, f"Expected 30, got {eligible_score_cell}"

    def test_sheet_1_signals_dash_when_no_eligible_score(self):
        import openpyxl
        data = self._make_data(None)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        build_sheet_1_signals(wb, data)
        ws = wb["1_Signals"]
        eligible_score_cell = ws.cell(row=2, column=14).value
        assert eligible_score_cell == "—", f"Expected dash, got {eligible_score_cell}"


# ─────────────────────────────────────────────────────────────────────────────
# Cost columns from DB (FIX-124 — replaces FIX-121 runtime computation)
# ─────────────────────────────────────────────────────────────────────────────

class TestCostColumnsFromDB:

    def test_build_sheet_2_orders_expense_cols_from_db(self):
        """Cols 31-35 (brokerage/STT/etc.) read from trades.cost_* DB columns."""
        import openpyxl
        trade = {
            "trade_id": "t-exp-01",
            "signal_id": "sig-exp-01",
            "symbol": "RELIANCE",
            "direction": "LONG",
            "strategy": "gap_go_long",
            "qty_planned": 10,
            "qty_filled": 10,
            "entry_target_price": 2500.0,
            "entry_actual_price": 2500.0,
            "exit_price": 2520.0,
            "sl_initial": 2450.0,
            "tgt_initial": 2600.0,
            "gross_pnl": 200.0,
            "charges": 55.0,
            "net_pnl": 145.0,
            "status": "CLOSED",
            "created_at": "2026-05-18T09:30:00+05:30",
            "entry_time": "2026-05-18T09:30:05+05:30",
            "exit_time": "2026-05-18T10:45:00+05:30",
            "exit_reason": "TGT_HIT",
            "cost_brokerage": 40.0,
            "cost_stt": 6.3,
            "cost_exchange_txn": 3.0,
            "cost_stamp_duty": 0.75,
            "cost_gst": 7.74,
            "sl_trail_count": 2,
        }
        data = ReportData(
            date_iso="2026-05-18",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=100145.0,
            signals=[],
            trades=[trade],
            orders=[],
            fm_ledger=[],
            screener_results=[],
            innings=[],
            system_events=[],
            recon_log=[],
            gate_state=[],
            excluded_symbols=[],
        )
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        build_sheet_2_orders(wb, data)
        ws = wb["2_Orders"]
        assert ws.cell(4, 31).value == 40.0
        assert ws.cell(4, 32).value == 6.3
        assert ws.cell(4, 33).value == 3.0
        assert ws.cell(4, 34).value == 0.75
        assert ws.cell(4, 35).value == 7.74
        assert ws.cell(4, 28).value == 2

    def test_build_sheet_2_orders_dash_when_no_cost_data(self):
        """Cols 31-35 show dash when trade has no cost columns."""
        import openpyxl
        trade = {
            "trade_id": "t-no-cost",
            "signal_id": "sig-nc",
            "symbol": "TCS",
            "direction": "LONG",
            "strategy": "gap_go_long",
            "qty_planned": 5,
            "qty_filled": 5,
            "entry_target_price": 3000.0,
            "entry_actual_price": 3000.0,
            "sl_initial": 2950.0,
            "tgt_initial": 3100.0,
            "gross_pnl": 0,
            "charges": 0,
            "net_pnl": 0,
            "status": "CLOSED",
            "created_at": "2026-05-18T09:30:00+05:30",
            "entry_time": "2026-05-18T09:30:05+05:30",
            "exit_time": "2026-05-18T10:45:00+05:30",
            "exit_reason": "EOD",
        }
        data = ReportData(
            date_iso="2026-05-18",
            mode="PAPER",
            account="TEST",
            opening_capital=100000.0,
            closing_capital_broker=100000.0,
            signals=[],
            trades=[trade],
            orders=[],
            fm_ledger=[],
            screener_results=[],
            innings=[],
            system_events=[],
            recon_log=[],
            gate_state=[],
            excluded_symbols=[],
        )
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        build_sheet_2_orders(wb, data)
        ws = wb["2_Orders"]
        for col in (31, 32, 33, 34, 35):
            assert ws.cell(4, col).value == "—"
