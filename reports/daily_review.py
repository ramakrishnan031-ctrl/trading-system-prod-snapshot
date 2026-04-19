"""
reports/daily_review.py — Trading System v2

Purpose:
    Generate EOD daily review report from state_store data.
    Produces .xlsx (primary) and/or .md (fallback/summary).
    Read-only: never mutates state_store.

Usage (CLI):
    python -m reports.daily_review [--date YYYY-MM-DD] [--output-dir DIR]
                                   [--format xlsx|md|both]

Usage (programmatic):
    generator = DailyReviewGenerator(state_store, time_authority, logger)
    paths = generator.generate("2026-04-16", Path("reports/daily"), ["xlsx", "md"])

Sections (9):
    1. SUMMARY            — totals, P&L, win rate, avg duration
    2. SIGNAL_FUNNEL      — stage-by-stage drop-off counts (P10)
    3. SCREENER_ANALYTICS — per-step rejections, tier dist, latencies (P18)
    4. TRADES             — per-trade row with P&L breakdown
    5. ORDERS             — per-order row with broker IDs and legs
    6. CAPITAL_LEDGER     — every fm_ledger entry today (BL-5 write-ahead log)
    7. SYSTEM_EVENTS      — kill-switch, reconciliation, lifecycle events
    8. ALERTS_SENT        — placeholder (no Telegram read-back in v2)
    9. MULTI_INNING_TRACKING — 36-col per-trade multi-inning simulation (DR-U1)

Locked: DR1-DR15, DR-U1-DR-U5
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

# HIGH #10: import openpyxl at module level so missing dep fails at startup
import openpyxl
from openpyxl.styles import Font

from core.time_authority import now_ist

_IST = timezone(timedelta(hours=5, minutes=30))


def _now_ist_date() -> str:
    return now_ist().strftime("%Y-%m-%d")


# ─────────────────────────────────────────────────────────────────────────────
# Output paths dataclass (DR6)
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class ReportPaths:
    xlsx_path: Optional[Path] = None
    md_path:   Optional[Path] = None


# ─────────────────────────────────────────────────────────────────────────────
# Signal funnel stages (P10)
# ─────────────────────────────────────────────────────────────────────────────

_FUNNEL_STAGE_STATUSES: Dict[str, List[str]] = {
    "RECEIVED":       [],                              # all signals
    "QUEUED":         ["QUEUED", "PROCESSING", "PASSED_SCREEN", "SIZED",
                       "APPROVED", "RESERVED", "PLACED", "FILLED",
                       "EXITED", "TRADED"],
    "PROCESSING":     ["PROCESSING", "PASSED_SCREEN", "SIZED", "APPROVED",
                       "RESERVED", "PLACED", "FILLED", "EXITED", "TRADED"],
    "PASSED_SCREEN":  ["PASSED_SCREEN", "SIZED", "APPROVED", "RESERVED",
                       "PLACED", "FILLED", "EXITED", "TRADED"],
    "SIZED":          ["SIZED", "APPROVED", "RESERVED", "PLACED",
                       "FILLED", "EXITED", "TRADED"],
    "APPROVED":       ["APPROVED", "RESERVED", "PLACED", "FILLED",
                       "EXITED", "TRADED"],
    "RESERVED":       ["RESERVED", "PLACED", "FILLED", "EXITED", "TRADED"],
    "PLACED":         ["PLACED", "FILLED", "EXITED", "TRADED"],
    "FILLED":         ["FILLED", "EXITED", "TRADED"],
    "EXITED":         ["EXITED", "TRADED"],
}


def _build_funnel(signals: List[dict]) -> List[dict]:
    """Build signal funnel rows from a list of signal dicts (DR9)."""
    total = len(signals)
    rows = []
    prev_count = total if total > 0 else 1   # avoid div-by-zero

    for stage, forward_statuses in _FUNNEL_STAGE_STATUSES.items():
        if stage == "RECEIVED":
            count = total
        else:
            count = sum(
                1 for s in signals
                if s.get("status") in forward_statuses
            )
        drop_pct = 0.0 if total == 0 else round((1 - count / prev_count) * 100, 1)
        rows.append({
            "stage":    stage,
            "count":    count,
            "drop_pct": drop_pct if stage != "RECEIVED" else 0.0,
        })
        prev_count = max(count, 1)

    return rows


# ─────────────────────────────────────────────────────────────────────────────
# Screener analytics (DR10)
# ─────────────────────────────────────────────────────────────────────────────

def _build_screener_analytics(screener_rows: List[dict]) -> Dict[str, Any]:
    """
    Aggregate screener_results rows:
      - per_step: {step_name: {count, avg_score}}
      - per_tier: {tier: {count, avg_score}}
      - latency:  {step_name: {avg_ms, p50_ms, p95_ms}}
    """
    per_step: Dict[str, dict] = {}
    per_tier: Dict[str, dict] = {}
    latencies: Dict[str, List[float]] = {}

    for row in screener_rows:
        status = row.get("status", "")
        tier   = row.get("tier", "UNKNOWN")
        score  = row.get("score", 0)

        # Per-tier
        if tier not in per_tier:
            per_tier[tier] = {"count": 0, "total_score": 0}
        per_tier[tier]["count"] += 1
        per_tier[tier]["total_score"] += score

        # Step results (JSON blob)
        try:
            step_results = json.loads(row.get("step_results", "{}"))
        except (json.JSONDecodeError, TypeError):
            step_results = {}

        for step_name, result in step_results.items():
            passed = result.get("passed", True)
            if not passed:
                if step_name not in per_step:
                    per_step[step_name] = {"count": 0, "total_score": 0}
                per_step[step_name]["count"] += 1
                per_step[step_name]["total_score"] += score

        # Latencies (JSON blob)
        try:
            lats = json.loads(row.get("latencies", "{}"))
        except (json.JSONDecodeError, TypeError):
            lats = {}

        for step_name, ms in lats.items():
            if isinstance(ms, (int, float)):
                latencies.setdefault(step_name, []).append(float(ms))

    # Compute averages
    per_step_out = {}
    for step, d in per_step.items():
        n = d["count"]
        per_step_out[step] = {
            "count":     n,
            "avg_score": round(d["total_score"] / n, 1) if n else 0.0,
        }

    per_tier_out = {}
    for tier, d in per_tier.items():
        n = d["count"]
        per_tier_out[tier] = {
            "count":     n,
            "avg_score": round(d["total_score"] / n, 1) if n else 0.0,
        }

    latency_out = {}
    for step, vals in latencies.items():
        vals.sort()
        n = len(vals)
        avg  = sum(vals) / n if n else 0.0
        p50  = vals[n // 2] if n else 0.0
        p95  = vals[int(n * 0.95)] if n >= 20 else (vals[-1] if vals else 0.0)
        latency_out[step] = {
            "avg_ms": round(avg, 2),
            "p50_ms": round(p50, 2),
            "p95_ms": round(p95, 2),
        }

    return {
        "per_step": per_step_out,
        "per_tier": per_tier_out,
        "latency":  latency_out,
    }


# ─────────────────────────────────────────────────────────────────────────────
# P&L summary (DR11)
# ─────────────────────────────────────────────────────────────────────────────

def _build_pnl_summary(trades: List[dict]) -> dict:
    """
    Compute daily P&L summary from trade rows (DR11).
    Uses trades.net_pnl directly (computed and written by fund_manager).
    """
    closed = [t for t in trades if t.get("status") == "CLOSED"]
    total_net   = sum(t.get("net_pnl") or 0.0 for t in closed)
    total_gross = sum(t.get("gross_pnl") or 0.0 for t in closed)
    total_cost  = sum(t.get("charges") or 0.0 for t in closed)
    winners     = [t for t in closed if (t.get("net_pnl") or 0.0) > 0]

    durations = []
    for t in closed:
        entry_t = t.get("entry_time")
        exit_t  = t.get("exit_time")
        if entry_t and exit_t:
            try:
                e = datetime.fromisoformat(entry_t)
                x = datetime.fromisoformat(exit_t)
                durations.append((x - e).total_seconds() / 60.0)
            except ValueError:
                pass

    avg_duration = round(sum(durations) / len(durations), 1) if durations else 0.0
    win_rate = round(len(winners) / len(closed) * 100, 1) if closed else 0.0

    return {
        "total_signals":       0,           # filled in by generate()
        "screened_passed":     0,
        "placed":              0,
        "completed":           len(closed),
        "realized_pnl":        round(total_net, 2),
        "gross_pnl":           round(total_gross, 2),
        "total_costs":         round(total_cost, 2),
        "win_rate_pct":        win_rate,
        "avg_duration_min":    avg_duration,
    }


# ─────────────────────────────────────────────────────────────────────────────
# Multi-inning tracking helpers (DR-U1, DR-U2)
# ─────────────────────────────────────────────────────────────────────────────

# 36-column header layout (DR-U2): 5 id + 9 per inning * 3 + 4 aggregate
_MULTI_INNING_HEADERS: List[str] = [
    # Identification (5)
    "trade_id", "signal_id", "symbol", "direction", "scanner_name",
    # Inning 1 (9)
    "i1_entry", "i1_entry_ts", "i1_sl", "i1_tgt",
    "i1_exit", "i1_exit_ts", "i1_exit_reason",
    "i1_pnl_pct", "i1_duration_min",
    # Inning 2 (9)
    "i2_entry", "i2_entry_ts", "i2_sl", "i2_tgt",
    "i2_exit", "i2_exit_ts", "i2_exit_reason",
    "i2_pnl_pct", "i2_duration_min",
    # Inning 3 (9)
    "i3_entry", "i3_entry_ts", "i3_sl", "i3_tgt",
    "i3_exit", "i3_exit_ts", "i3_exit_reason",
    "i3_pnl_pct", "i3_duration_min",
    # Aggregate (4)
    "total_innings", "cumulative_simulated_pnl_pct",
    "best_inning_pnl_pct", "worst_inning_pnl_pct",
]


def _fmt_ts(ts_str: Optional[str]) -> str:
    """Normalise an ISO timestamp to 'YYYY-MM-DD HH:MM:SS' or empty string."""
    if not ts_str:
        return ""
    try:
        dt = datetime.fromisoformat(ts_str)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        return str(ts_str)


def _pivot_innings_to_trade_rows(flat_rows: List[dict]) -> List[dict]:
    """
    Pivot flat per-inning rows into one dict per trade (DR-U2).

    Input:  list of dicts from get_inning_summary_by_date (ordered by trade_id,
            inning_number).
    Output: list of dicts with exactly 36 columns each.
            i2/i3 blocks are blank strings when trade has fewer innings.
            exit_reason is "OPEN" when no exit recorded.
    """
    # Group by trade_id (preserving first-seen order)
    grouped: Dict[str, List[dict]] = {}
    for row in flat_rows:
        tid = row["trade_id"]
        grouped.setdefault(tid, []).append(row)

    result = []
    for tid, innings in grouped.items():
        # Sort by inning_number to ensure correct assignment
        innings_sorted = sorted(innings, key=lambda r: r["inning_number"])

        # Map inning_number -> row
        by_num: Dict[int, dict] = {r["inning_number"]: r for r in innings_sorted}

        # Identity fields from inning 1 (always present)
        i1 = by_num.get(1, innings_sorted[0])
        row_out: Dict[str, Any] = {
            "trade_id":     tid,
            "signal_id":    i1.get("signal_id", ""),
            "symbol":       i1.get("symbol", ""),
            "direction":    i1.get("direction", ""),
            "scanner_name": i1.get("scanner_name", ""),
        }

        # Inning blocks
        all_pnls: List[float] = []
        sim_pnls: List[float] = []
        for n in (1, 2, 3):
            pfx = f"i{n}_"
            ing = by_num.get(n)
            if ing is None:
                for col in ("entry", "entry_ts", "sl", "tgt", "exit",
                            "exit_ts", "exit_reason", "pnl_pct", "duration_min"):
                    row_out[pfx + col] = ""
            else:
                row_out[pfx + "entry"]      = ing.get("entry_price", "")
                row_out[pfx + "entry_ts"]   = _fmt_ts(ing.get("entry_ts"))
                row_out[pfx + "sl"]         = ing.get("sl_price", "")
                row_out[pfx + "tgt"]        = ing.get("tgt_price", "")
                # Exit fields — blank if OPEN (no exit yet)
                has_exit = ing.get("exit_price") is not None
                row_out[pfx + "exit"]       = ing.get("exit_price", "") if has_exit else ""
                row_out[pfx + "exit_ts"]    = _fmt_ts(ing.get("exit_ts")) if has_exit else ""
                row_out[pfx + "exit_reason"] = (
                    ing.get("exit_reason") or "OPEN"
                    if ing.get("exit_reason")
                    else "OPEN"
                )
                pnl = ing.get("pnl_pct")
                row_out[pfx + "pnl_pct"]    = round(pnl, 2) if pnl is not None and has_exit else ""
                dur = ing.get("duration_sec")
                row_out[pfx + "duration_min"] = int(dur // 60) if dur is not None and has_exit else ""
                # Collect for aggregates
                if pnl is not None and has_exit:
                    all_pnls.append(pnl)
                    if n > 1:
                        sim_pnls.append(pnl)

        # Aggregate block
        row_out["total_innings"] = len(by_num)
        row_out["cumulative_simulated_pnl_pct"] = (
            round(sum(sim_pnls), 2) if sim_pnls else ""
        )
        row_out["best_inning_pnl_pct"]  = round(max(all_pnls), 2) if all_pnls else ""
        row_out["worst_inning_pnl_pct"] = round(min(all_pnls), 2) if all_pnls else ""

        result.append(row_out)
    return result


def _build_multi_inning_summary(trade_rows: List[dict]) -> dict:
    """
    Compute the multi-inning sub-block for the SUMMARY section (DR-U4).

    Returns a dict of label -> value for each stat line.
    """
    n_trades = len(trade_rows)
    n_i1_tgt  = sum(1 for r in trade_rows if r.get("i1_exit_reason") == "TGT")
    n_i2      = sum(1 for r in trade_rows if r.get("i2_entry") != "")
    n_i2_tgt  = sum(1 for r in trade_rows if r.get("i2_exit_reason") == "TGT")
    n_i3      = sum(1 for r in trade_rows if r.get("i3_entry") != "")
    n_i3_tgt  = sum(1 for r in trade_rows if r.get("i3_exit_reason") == "TGT")
    sim_pnls  = [
        r["cumulative_simulated_pnl_pct"]
        for r in trade_rows
        if isinstance(r.get("cumulative_simulated_pnl_pct"), float)
    ]
    cum_sim_pnl = round(sum(sim_pnls), 2) if sim_pnls else 0.0

    return {
        "Trades with inning tracking":          n_trades,
        "Trades with inning 1 TGT hit":         n_i1_tgt,
        "Trades with inning 2 reached":         n_i2,
        "Trades with inning 2 TGT hit (sim)":   n_i2_tgt,
        "Trades with inning 3 reached":         n_i3,
        "Trades with inning 3 TGT hit (sim)":   n_i3_tgt,
        "Cumulative simulated P&L (innings 2+3)": cum_sim_pnl,
        "Note":                                 (
            "Simulated innings are hypothetical tracking only. "
            "No real orders placed for innings 2 and 3."
        ),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Report builder
# ─────────────────────────────────────────────────────────────────────────────

class DailyReviewGenerator:
    """
    Generates daily review reports from state_store data (DR6).
    Read-only: never mutates state_store (DR7).
    """

    def __init__(self, state_store, time_authority, logger: logging.Logger):
        self._store         = state_store
        self._time_authority = time_authority
        self._log           = logger

    # ── public API ────────────────────────────────────────────────────────────

    def generate(
        self,
        date_iso: str,
        output_dir: Path,
        formats: List[str],
    ) -> ReportPaths:
        """
        Generate report for date_iso.
        formats: subset of ["xlsx", "md"].
        Creates output_dir if absent (DR12, DR4).
        """
        self._log.info("daily_review: generating for %s formats=%s", date_iso, formats)

        # Collect data
        data = self._collect(date_iso)

        # Ensure output dir exists (DR4)
        day_dir = output_dir / date_iso
        day_dir.mkdir(parents=True, exist_ok=True)

        paths = ReportPaths()

        if "xlsx" in formats:
            paths.xlsx_path = self._write_xlsx(data, day_dir)
            self._log.info("daily_review: xlsx -> %s", paths.xlsx_path)

        if "md" in formats:
            paths.md_path = self._write_md(data, day_dir)
            self._log.info("daily_review: md -> %s", paths.md_path)

        return paths

    # ── data collection ───────────────────────────────────────────────────────

    def _collect(self, date_iso: str) -> dict:
        """Query all data sources for the given date (DR8). Read-only."""
        signals          = self._store.get_signals_for_date(date_iso)
        trades           = self._store.get_trades_for_date(date_iso)
        orders           = self._store.get_orders_for_date(date_iso)
        fm_ledger        = self._store.get_fm_ledger_for_date(date_iso)
        system_events    = self._store.get_system_events_for_date(date_iso)
        recon_log        = self._store.get_reconciliation_log_for_date(date_iso)
        screener_results = self._store.get_screener_results_for_date(date_iso)

        # Multi-inning tracking (DR-U3)
        inning_flat  = self._store.get_inning_summary_by_date(date_iso)
        inning_rows  = _pivot_innings_to_trade_rows(inning_flat)
        mi_summary   = _build_multi_inning_summary(inning_rows)

        funnel    = _build_funnel(signals)
        analytics = _build_screener_analytics(screener_results)
        summary   = _build_pnl_summary(trades)

        # Fill in signal-level counters in summary
        summary["total_signals"]   = len(signals)
        summary["screened_passed"] = sum(
            1 for s in signals
            if (s.get("status") or "").startswith("PASSED")
            or s.get("status") in ("TRADED", "PLACED", "FILLED", "EXITED",
                                    "SIZED", "APPROVED", "RESERVED")
        )
        summary["placed"] = sum(
            1 for s in signals
            if s.get("status") in ("PLACED", "FILLED", "EXITED", "TRADED")
        )

        return {
            "date_iso":        date_iso,
            "summary":         summary,
            "funnel":          funnel,
            "analytics":       analytics,
            "trades":          trades,
            "orders":          orders,
            "fm_ledger":       fm_ledger,
            "system_events":   system_events,
            "recon_log":       recon_log,
            "screener_results": screener_results,
            "inning_rows":     inning_rows,
            "mi_summary":      mi_summary,
        }

    # ── XLSX writer ───────────────────────────────────────────────────────────

    def _write_xlsx(self, data: dict, out_dir: Path) -> Path:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)   # remove default sheet

        self._xlsx_summary(wb, data)
        self._xlsx_funnel(wb, data)
        self._xlsx_screener(wb, data)
        self._xlsx_trades(wb, data)
        self._xlsx_orders(wb, data)
        self._xlsx_capital(wb, data)
        self._xlsx_system(wb, data)
        self._xlsx_alerts(wb, data)
        self._xlsx_multi_inning(wb, data)

        path = out_dir / "daily_review.xlsx"
        wb.save(path)
        return path

    def _xlsx_sheet(self, wb, title: str, headers: List[str], rows: List[List]):
        """Helper: create a sheet with bold header row."""
        ws = wb.create_sheet(title=title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        return ws

    def _xlsx_summary(self, wb, data: dict) -> None:
        s = data["summary"]
        rows = [[k, v] for k, v in s.items()]
        # Multi-inning sub-block (DR-U4)
        rows.append(["--- Multi-Inning Tracking ---", ""])
        for k, v in data["mi_summary"].items():
            rows.append([k, v])
        self._xlsx_sheet(wb, "SUMMARY", ["metric", "value"], rows)

    def _xlsx_funnel(self, wb, data: dict) -> None:
        rows = [
            [f["stage"], f["count"], f["drop_pct"]]
            for f in data["funnel"]
        ]
        self._xlsx_sheet(wb, "SIGNAL_FUNNEL",
                         ["stage", "count", "drop_pct_%"], rows)

    def _xlsx_screener(self, wb, data: dict) -> None:
        analytics = data["analytics"]
        rows = []
        for step, d in analytics["per_step"].items():
            rows.append(["step_rejection", step, d["count"], d["avg_score"], "", ""])
        for tier, d in analytics["per_tier"].items():
            rows.append(["tier_dist", tier, d["count"], d["avg_score"], "", ""])
        for step, d in analytics["latency"].items():
            rows.append(["latency", step, "", "",
                         d["avg_ms"], d["p95_ms"]])
        self._xlsx_sheet(wb, "SCREENER_ANALYTICS",
                         ["type", "name", "count", "avg_score",
                          "avg_ms", "p95_ms"], rows)

    def _xlsx_trades(self, wb, data: dict) -> None:
        headers = [
            "trade_id", "signal_id", "symbol", "direction", "strategy",
            "qty_planned", "qty_filled", "entry_target_price", "entry_actual_price",
            "sl_initial", "tgt_initial", "gross_pnl", "charges", "net_pnl",
            "created_at", "entry_time", "exit_time", "exit_reason", "status",
            "order_protocol",
        ]
        rows = [[t.get(h) for h in headers] for t in data["trades"]]
        self._xlsx_sheet(wb, "TRADES", headers, rows)

    def _xlsx_orders(self, wb, data: dict) -> None:
        headers = [
            "order_id", "trade_id", "leg", "transaction_type", "order_type",
            "product", "qty_requested", "price", "trigger_price",
            "status", "qty_filled", "avg_fill_price",
            "placed_at", "updated_at",
        ]
        rows = [[o.get(h) for h in headers] for o in data["orders"]]
        self._xlsx_sheet(wb, "ORDERS", headers, rows)

    def _xlsx_capital(self, wb, data: dict) -> None:
        # BL-5: fm_ledger extended with entry_type (was mutation_type),
        # session_id, direction, trade_id, margin_delta, pnl_delta, costs.
        headers = [
            "ledger_id", "ts", "entry_type", "amount", "bucket",
            "balance_before", "balance_after",
            "signal_id", "reservation_id", "reason",
            "session_id", "direction", "trade_id",
            "margin_delta", "pnl_delta", "costs",
        ]
        rows = [[r.get(h) for h in headers] for r in data["fm_ledger"]]
        self._xlsx_sheet(wb, "CAPITAL_LEDGER", headers, rows)

    def _xlsx_system(self, wb, data: dict) -> None:
        # Merge system_events and reconciliation_log
        headers = ["source", "ts", "type", "tier", "symbol",
                   "description", "action_taken", "success", "details"]
        rows = []
        for e in data["system_events"]:
            rows.append([
                "system_event",
                e.get("timestamp"), e.get("event_type"),
                "", "", e.get("scenario", ""), "", "", e.get("details", ""),
            ])
        for r in data["recon_log"]:
            rows.append([
                "recon_log",
                r.get("ts"), r.get("check_name"),
                r.get("tier"), r.get("symbol"),
                r.get("description"), r.get("action_taken"),
                r.get("success"), "",
            ])
        self._xlsx_sheet(wb, "SYSTEM_EVENTS", headers, rows)

    def _xlsx_alerts(self, wb, data: dict) -> None:
        # Placeholder: no Telegram read-back in v2 (DR3 section 8)
        self._xlsx_sheet(wb, "ALERTS_SENT",
                         ["note"],
                         [["Telegram alert history not stored in state_store v2."]])

    def _xlsx_multi_inning(self, wb, data: dict) -> None:
        """Write MULTI_INNING_TRACKING sheet (DR-U5)."""
        inning_rows = data["inning_rows"]
        ws = wb.create_sheet(title="MULTI_INNING_TRACKING")
        ws.append(_MULTI_INNING_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"

        if not inning_rows:
            ws.append(["No multi-inning data for this date"] + [""] * (len(_MULTI_INNING_HEADERS) - 1))
        else:
            for row_dict in inning_rows:
                ws.append([row_dict.get(h, "") for h in _MULTI_INNING_HEADERS])

        # Auto-size columns (min 10, max 25)
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) if cell.value is not None else 0)
                for cell in col
            )
            ws.column_dimensions[col[0].column_letter].width = max(10, min(25, max_len + 2))

    # ── Markdown writer ───────────────────────────────────────────────────────

    def _write_md(self, data: dict, out_dir: Path) -> Path:
        lines: List[str] = []
        d = data["date_iso"]
        lines.append(f"# Daily Review — {d}\n")

        self._md_summary(lines, data)
        self._md_funnel(lines, data)
        self._md_screener(lines, data)
        self._md_trades(lines, data)
        self._md_orders(lines, data)
        self._md_capital(lines, data)
        self._md_system(lines, data)
        self._md_alerts(lines, data)
        self._md_multi_inning(lines, data)

        path = out_dir / "daily_review.md"
        path.write_text("\n".join(lines), encoding="utf-8")
        return path

    def _md_table(self, lines: List[str], headers: List[str],
                  rows: List[List]) -> None:
        lines.append("| " + " | ".join(str(h) for h in headers) + " |")
        lines.append("| " + " | ".join("---" for _ in headers) + " |")
        if not rows:
            lines.append("| " + " | ".join("—" for _ in headers) + " |")
        else:
            for row in rows:
                cells = [str(v) if v is not None else "" for v in row]
                lines.append("| " + " | ".join(cells) + " |")
        lines.append("")

    def _md_summary(self, lines: List[str], data: dict) -> None:
        s = data["summary"]
        lines.append("## 1. Summary\n")
        if not data["trades"] and not data["summary"]["total_signals"]:
            lines.append("*No activity today.*\n")
        # Merge main summary + multi-inning sub-block (DR-U4)
        all_rows = list(s.items())
        all_rows.append(("--- Multi-Inning Tracking ---", ""))
        all_rows.extend(data["mi_summary"].items())
        self._md_table(lines, ["metric", "value"], all_rows)

    def _md_funnel(self, lines: List[str], data: dict) -> None:
        lines.append("## 2. Signal Funnel\n")
        rows = [[f["stage"], f["count"], f["drop_pct"]] for f in data["funnel"]]
        self._md_table(lines, ["stage", "count", "drop_pct_%"], rows)

    def _md_screener(self, lines: List[str], data: dict) -> None:
        lines.append("## 3. Screener Analytics\n")
        analytics = data["analytics"]
        rows = []
        for step, d in analytics["per_step"].items():
            rows.append(["rejection", step, d["count"], d["avg_score"], "", ""])
        for tier, d in analytics["per_tier"].items():
            rows.append(["tier", tier, d["count"], d["avg_score"], "", ""])
        for step, d in analytics["latency"].items():
            rows.append(["latency", step, "", "", d["avg_ms"], d["p95_ms"]])
        self._md_table(lines, ["type", "name", "count", "avg_score",
                                "avg_ms", "p95_ms"], rows)

    def _md_trades(self, lines: List[str], data: dict) -> None:
        lines.append("## 4. Trades\n")
        headers = ["trade_id", "symbol", "direction", "qty_filled",
                   "gross_pnl", "charges", "net_pnl", "status", "exit_reason"]
        rows = [[t.get(h) for h in headers] for t in data["trades"]]
        self._md_table(lines, headers, rows)

    def _md_orders(self, lines: List[str], data: dict) -> None:
        lines.append("## 5. Orders\n")
        headers = ["order_id", "trade_id", "leg", "status",
                   "qty_requested", "price", "placed_at"]
        rows = [[o.get(h) for h in headers] for o in data["orders"]]
        self._md_table(lines, headers, rows)

    def _md_capital(self, lines: List[str], data: dict) -> None:
        lines.append("## 6. Capital Ledger\n")
        headers = ["ts", "entry_type", "amount", "bucket",
                   "balance_before", "balance_after", "reservation_id"]
        rows = [[r.get(h) for h in headers] for r in data["fm_ledger"]]
        self._md_table(lines, headers, rows)

    def _md_system(self, lines: List[str], data: dict) -> None:
        lines.append("## 7. System Events & Reconciliation\n")
        headers = ["source", "ts", "type", "details"]
        rows = []
        for e in data["system_events"]:
            rows.append(["system", e.get("timestamp"), e.get("event_type"),
                         e.get("details") or ""])
        for r in data["recon_log"]:
            rows.append(["recon", r.get("ts"), r.get("check_name"),
                         r.get("description") or ""])
        self._md_table(lines, headers, rows)

    def _md_alerts(self, lines: List[str], data: dict) -> None:
        lines.append("## 8. Alerts Sent\n")
        lines.append("*Telegram alert history not stored in state_store v2.*\n")

    def _md_multi_inning(self, lines: List[str], data: dict) -> None:
        """Write Multi-Inning Tracking section (DR-U5)."""
        lines.append("## Multi-Inning Tracking\n")
        inning_rows = data["inning_rows"]
        if not inning_rows:
            lines.append("*No multi-inning data for this date.*\n")
            return
        rows = [[r.get(h, "") for h in _MULTI_INNING_HEADERS] for r in inning_rows]
        self._md_table(lines, _MULTI_INNING_HEADERS, rows)


# ─────────────────────────────────────────────────────────────────────────────
# CLI entry point (DR2)
# ─────────────────────────────────────────────────────────────────────────────

def _parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="reports.daily_review",
        description="Generate EOD daily review report (DR1-DR15).",
    )
    parser.add_argument(
        "--date",
        metavar="YYYY-MM-DD",
        default=None,
        help="Date to generate report for (default: today IST).",
    )
    parser.add_argument(
        "--output-dir",
        metavar="DIR",
        default="reports/daily",
        dest="output_dir",
        help="Output directory (default: reports/daily/).",
    )
    parser.add_argument(
        "--format",
        choices=["xlsx", "md", "both"],
        default="both",
        dest="fmt",
        help="Output format (default: both).",
    )
    parser.add_argument(
        "--db",
        metavar="PATH",
        default="data_store/trading_system.db",
        help="Path to SQLite database (default: data_store/trading_system.db).",
    )
    parser.add_argument(
        "--unattended",
        action="store_true",
        help="Suppress stdout decoration (for systemd journal / cron).",
    )
    return parser.parse_args(argv)


def main(argv=None) -> int:
    args = _parse_args(argv)

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s %(message)s",
    )
    log = logging.getLogger("daily_review")

    date_iso = args.date or _now_ist_date()

    # Validate date format
    try:
        date.fromisoformat(date_iso)
    except ValueError:
        print(f"ERROR: invalid date: {date_iso!r} — expected YYYY-MM-DD",
              file=sys.stderr)
        return 1

    from core.state_store import StateStore
    import core.time_authority as time_authority

    db_path = Path(args.db)
    if not db_path.exists():
        print(f"ERROR: database not found: {db_path}", file=sys.stderr)
        return 1

    store = StateStore(db_path)

    formats = ["xlsx", "md"] if args.fmt == "both" else [args.fmt]
    generator = DailyReviewGenerator(store, time_authority, log)
    try:
        paths = generator.generate(date_iso, Path(args.output_dir), formats)
    except Exception as exc:
        print(
            f"ERROR: daily_review.generate failed: {exc}",
            file=sys.stderr,
        )
        log.exception("daily_review.generate failed")
        return 2

    if not args.unattended:
        if paths.xlsx_path:
            print(f"xlsx: {paths.xlsx_path}")
        if paths.md_path:
            print(f"md:   {paths.md_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
